import io
import json
import os
import re
from collections import Counter
from typing import List, Dict, Any

import streamlit as st
from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document as DocxDocument

from langchain_core.documents import Document
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_text_splitters import RecursiveCharacterTextSplitter

load_dotenv()

st.set_page_config(page_title="Interview Question Generator", page_icon="🎯", layout="wide")

DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")


def extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if name.endswith(".docx"):
        doc = DocxDocument(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if name.endswith((".txt", ".md", ".markdown")):
        return data.decode("utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {uploaded_file.name}")


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def make_documents(files) -> List[Document]:
    docs = []
    for file in files:
        raw = clean_text(extract_text(file))
        if not raw:
            st.warning(f"No extractable text found in {file.name}. Scanned PDFs may need OCR first.")
            continue
        docs.append(Document(page_content=raw, metadata={"source": file.name}))
    if not docs:
        raise ValueError("No readable text was found in the uploaded files.")
    splitter = RecursiveCharacterTextSplitter(chunk_size=700, chunk_overlap=80, add_start_index=True)
    chunks = splitter.split_documents(docs)
    # Bound memory use on small hosting instances.
    if len(chunks) > 300:
        chunks = chunks[:300]
        st.warning("The document set was large, so only the first 300 text chunks were processed. Upload smaller files for better coverage.")
    for i, chunk in enumerate(chunks):
        chunk.metadata["chunk_id"] = i
    return chunks


def get_llm(api_key: str, model_name: str):
    if not api_key.strip():
        raise ValueError("Enter your Gemini API key in the sidebar or set GEMINI_API_KEY in .env.")
    return ChatGoogleGenerativeAI(
        model=model_name,
        google_api_key=api_key.strip(),
        temperature=0.25,
        max_retries=2,
    )


def invoke_text(llm, prompt: str) -> str:
    response = llm.invoke(prompt)
    return str(response.content).strip()


def retrieve_context(chunks, topic: str, k: int):
    """Lightweight lexical retrieval; avoids loading a neural embedding model or vector DB."""
    stop_words = {
        "the", "and", "for", "with", "that", "this", "from", "into", "are", "was",
        "were", "you", "your", "what", "when", "where", "which", "how", "why",
        "about", "using", "use", "can", "could", "should", "would", "have", "has",
        "had", "not", "but", "all", "any", "its", "their", "they", "them", "then",
        "than", "also", "each", "such", "these", "those", "our", "out", "more"
    }
    query = (topic or "").lower()
    terms = [t for t in re.findall(r"[a-zA-Z0-9_+#.-]{2,}", query) if t not in stop_words]
    if not terms:
        return chunks[:min(k, len(chunks))]
    scored = []
    for index, chunk in enumerate(chunks):
        text = chunk.page_content.lower()
        words = re.findall(r"[a-zA-Z0-9_+#.-]{2,}", text)
        counts = Counter(words)
        score = sum(min(counts.get(term, 0), 4) for term in terms)
        # Stable tie-break preserves document order.
        scored.append((score, -index, chunk))
    scored.sort(reverse=True, key=lambda item: (item[0], item[1]))
    matches = [item[2] for item in scored[:k] if item[0] > 0]
    return matches or chunks[:min(k, len(chunks))]


def generate_questions(llm, chunks, topic: str, difficulty: str, count: int, role: str) -> List[Dict[str, str]]:
    context = "\n\n".join(
        f"[Source: {d.metadata.get('source', 'document')}, chunk {d.metadata.get('chunk_id', '?')}]\n{d.page_content}"
        for d in chunks
    )
    prompt = f"""
You are an interview-question designer. Use the provided source context as the primary source.
Create exactly {count} distinct interview questions.
Topic: {topic or 'Choose the most relevant topics in the material'}
Target role: {role or 'General technical interview'}
Difficulty: {difficulty}

Rules:
- Questions must be answerable from the context or be clearly inferable from it.
- Avoid duplicates and vague questions.
- Cover different concepts when the context allows.
- Do not invent claims about the source.
- Return ONLY valid JSON: an array of objects with keys "question", "concept", and "source_hint".
- source_hint should be a short quote or description of the relevant context.

SOURCE CONTEXT:
{context}
"""
    raw = invoke_text(llm, prompt)
    # Tolerate markdown fences or small surrounding text.
    match = re.search(r"\[.*\]", raw, flags=re.S)
    if not match:
        raise ValueError("Gemini did not return a JSON array. Try fewer questions or a shorter topic.")
    parsed = json.loads(match.group(0))
    if not isinstance(parsed, list):
        raise ValueError("Unexpected output format from Gemini.")
    results = []
    for item in parsed:
        if isinstance(item, dict) and item.get("question"):
            results.append({
                "question": str(item["question"]).strip(),
                "concept": str(item.get("concept", "")).strip(),
                "source_hint": str(item.get("source_hint", "")).strip(),
            })
    if not results:
        raise ValueError("No usable questions were generated.")
    return results[:count]


def refine_question(llm, question: Dict[str, str], chunks, topic: str) -> Dict[str, str]:
    # Iterative refinement: update the draft using one retrieved chunk at a time.
    current = question["question"]
    used = []
    for chunk in chunks:
        prompt = f"""
Refine this technical interview question using the additional source passage.
Keep it clear, answerable, and focused. Do not make it needlessly complex.
If the passage adds no useful information, keep the question essentially unchanged.
Topic: {topic or 'General'}
Current question: {current}
Additional passage:
{chunk.page_content}
Return only the revised question, with no commentary.
"""
        current = invoke_text(llm, prompt).strip().strip('"')
        used.append({
            "source": chunk.metadata.get("source", "document"),
            "chunk_id": chunk.metadata.get("chunk_id", "?"),
            "text": chunk.page_content,
        })
    question["question"] = current
    question["sources"] = used
    return question


def generate_answer(llm, question: str, chunks) -> str:
    context = "\n\n".join(
        f"[{d.metadata.get('source', 'document')} | chunk {d.metadata.get('chunk_id', '?')}]\n{d.page_content}"
        for d in chunks
    )
    prompt = f"""
Answer this interview question using the supplied context.
Question: {question}

Requirements:
- Give a clear, technically accurate, interview-ready answer.
- Use the context as the source of truth.
- If the context does not contain enough information, explicitly say what is missing rather than inventing details.
- Use a short example only when supported or when it is a general illustrative example clearly labeled as such.
- Structure the answer with concise paragraphs or bullets.

CONTEXT:
{context}
"""
    return invoke_text(llm, prompt)


def main():
    st.title("🎯 LLM-Based Interview Question Generator")
    st.caption("Generate document-grounded interview questions and answers with RAG + Gemini. No account or login required.")

    with st.sidebar:
        st.header("Configuration")
        env_key = os.getenv("GEMINI_API_KEY", "")
        api_key = st.text_input("Gemini API key", value=env_key, type="password", help="Prefer .env for local use. Never commit your API key.")
        model_name = st.text_input("Gemini model", value=DEFAULT_MODEL)
        st.divider()
        st.caption("Lightweight local text retrieval is used; Gemini API generates questions and answers.")

    left, right = st.columns([1, 1])
    with left:
        st.subheader("1. Add source material")
        files = st.file_uploader(
            "Upload coding notes or documentation",
            type=["pdf", "docx", "txt", "md", "markdown"],
            accept_multiple_files=True,
        )
        st.caption("Text-based PDFs are supported. Scanned/image-only PDFs need OCR before upload.")
    with right:
        st.subheader("2. Configure questions")
        topic = st.text_input("Topic or focus", placeholder="e.g., Python, operating systems, RAG")
        role = st.text_input("Target role", placeholder="e.g., Python developer, ML engineer")
        difficulty = st.select_slider("Difficulty", options=["Beginner", "Intermediate", "Advanced"], value="Intermediate")
        count = st.slider("Number of questions", min_value=1, max_value=10, value=5)
        retrieval_k = st.slider("Context chunks to retrieve", min_value=2, max_value=10, value=5)

    if st.button("Generate interview questions + answers", type="primary", use_container_width=True):
        if not files:
            st.error("Upload at least one source document first.")
            return
        try:
            with st.spinner("Extracting text and splitting documents into chunks..."):
                chunks = make_documents(files)
            st.info(f"Prepared {len(chunks)} text chunks from {len(files)} file(s).")
            with st.spinner("Selecting relevant text and generating questions with Gemini..."):
                retrieved = retrieve_context(chunks, topic, retrieval_k)
                llm = get_llm(api_key, model_name)
                questions = generate_questions(llm, retrieved, topic, difficulty, count, role)
            with st.spinner("Generating interview-ready answers..."):
                results = []
                for item in questions:
                    answer = generate_answer(llm, item["question"], retrieved)
                    item["sources"] = [
                        {
                            "source": d.metadata.get("source", "document"),
                            "chunk_id": d.metadata.get("chunk_id", "?"),
                            "text": d.page_content,
                        }
                        for d in retrieved
                    ]
                    results.append({**item, "answer": answer})
            st.session_state["results"] = results
            st.session_state["run_meta"] = {
                "files": [f.name for f in files],
                "chunks": len(chunks),
                "model": model_name,
                "topic": topic,
                "difficulty": difficulty,
            }
            st.success(f"Generated {len(results)} question-and-answer pair(s).")
        except Exception as exc:
            st.error(f"Could not complete generation: {exc}")
            st.caption("Check the Gemini API key/model, network access, and that uploaded documents contain extractable text. If the host restarts, reduce document size.")

    if "results" in st.session_state:
        st.divider()
        st.subheader("Generated interview set")
        meta = st.session_state.get("run_meta", {})
        st.caption(f"Sources: {', '.join(meta.get('files', []))} · Chunks: {meta.get('chunks', 0)} · Model: {meta.get('model', '')}")
        export = []
        for idx, item in enumerate(st.session_state["results"], start=1):
            with st.expander(f"Q{idx}. {item['question']}", expanded=(idx == 1)):
                if item.get("concept"):
                    st.markdown(f"**Concept:** {item['concept']}")
                st.markdown("**Answer**")
                st.write(item["answer"])
                if item.get("source_hint"):
                    st.markdown("**Generated source hint**")
                    st.caption(item["source_hint"])
                sources = item.get("sources", [])
                if sources:
                    st.markdown("**Retrieved source passages**")
                    seen = set()
                    for source in sources:
                        key = (source["source"], source["chunk_id"])
                        if key in seen:
                            continue
                        seen.add(key)
                        with st.expander(f"{source['source']} — chunk {source['chunk_id']}"):
                            st.write(source["text"])
            export.append(
                f"## Question {idx}\n{item['question']}\n\n### Answer\n{item['answer']}\n"
                + (f"\nConcept: {item['concept']}\n" if item.get("concept") else "")
            )
        st.download_button(
            "Download questions and answers (.md)",
            data="\n\n---\n\n".join(export),
            file_name="interview_questions_and_answers.md",
            mime="text/markdown",
            use_container_width=True,
        )
        if st.button("Clear generated results"):
            st.session_state.pop("results", None)
            st.session_state.pop("run_meta", None)
            st.rerun()

    st.divider()
    st.caption("Privacy note: there is no login system. Uploaded documents are processed for the current session; generated prompts/context are sent to Gemini for generation. Do not upload confidential material unless you are comfortable sending it to the API provider.")


if __name__ == "__main__":
    main()
